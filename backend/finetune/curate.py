"""Phase 2 — fine-tuning data curation pipeline.

Turns a directory of raw broker documents (PDF / Word / CSV / text) into a
versioned SFT dataset. A strong model (Claude) does the labelling for you:
per document it classifies, quality-gates, reconstructs the training *input*,
preserves the human *output*, and redacts PII. You do NOT hand-label — you skim
the report + samples before training. See README.md for the full flow.

Design notes (why it's built this way) live in `MIGRATION.md` §4.7–§4.8 + D13/D19.

Run a real curation pass:
    cd backend
    CURATE_API_KEY=sk-ant-...  python -m finetune.curate \
        --input-dir /path/to/docs --out-dir finetune/dataset

Offline structural self-check (no API, no docs):
    python -m finetune.curate --selftest
"""

from __future__ import annotations

import argparse
import asyncio
import json
import os
import random
from pathlib import Path

# Memos are a few pages; cap the text we send so curation token cost stays bounded.
MAX_DOC_CHARS = 40_000
SUPPORTED_EXT = {".pdf", ".docx", ".txt", ".md", ".csv", ".tsv"}

# The curation engine prompt. Redaction is mandatory (owner decision: redact all
# PII / client confidentiality). The model preserves the human analysis verbatim
# (lightly cleaned) — that human voice + judgment is the asset we're training on.
CURATION_SYSTEM = """You are a data-curation engine building fine-tuning examples for Space Goose, an AI commercial-real-estate advisor. Goal: turn ONE broker document into a single training example that teaches a model to read a property + data, judge fit for a specific client's book of business, and give a clear, appropriately-skeptical go/no-go in a broker's voice.

You are given the full text of one document. Do four things:

1. CLASSIFY: is it a `void_analysis`, an `investment_memo`, or `other`? Only the first two are in scope.
2. QUALITY-GATE. Keep it ONLY if it (a) is a finished, coherent analysis (not a draft, template, or fragment), (b) reasons about a specific property and reaches a recommendation/assessment, and (c) shows real judgment — weighs pros/cons, flags risks, isn't boilerplate. Be strict: a fine-tune amplifies its data, so reject the mediocre.
3. If usable, BUILD the training pair:
   - input: reconstruct the request the advisor was answering — the property, the key data it had (demographics, traffic, tenants, financials), the client / book-of-business context if stated or inferable, and the question. Use ONLY facts present in the document; invent nothing. Write it as a realistic user request to a CRE assistant.
   - output: the document's analysis + recommendation, lightly normalized — lead with the go/no-go, strip letterhead/headers/footers/formatting cruft. DO NOT rewrite, improve, or restyle the analysis; preserve the human judgment and voice exactly.
4. REDACT (mandatory): replace every client name, contact, specific company identity, exact street address, and other PII in BOTH input and output with neutral placeholders (e.g. "the client", "[Tenant A]", "a ~96k-population trade area"). List what you redacted.

Return ONLY a JSON object, no prose, no code fences:
{
  "doc_type": "void_analysis" | "investment_memo" | "other",
  "usable": true | false,
  "quality": "high" | "medium" | "low",
  "reject_reason": "string (empty if usable)",
  "had_source_data": true | false,
  "pii_redacted": ["..."],
  "input": "string (empty if unusable)",
  "output": "string (empty if unusable)"
}"""


def advisory_system() -> str:
    """The system prompt the SFT records train under — reused from the eval so
    training, eval, and serving all share one advisory persona (train/serve
    consistency; see Gap 2). Phase 2 reconciles this with the production
    synthesis prompt."""
    from evals.judge import ADVISORY_SYSTEM

    return ADVISORY_SYSTEM


def extract_text(path: Path) -> str:
    """Best-effort text extraction by extension. Returns '' on any failure
    (the doc is skipped, never crashes the run)."""
    ext = path.suffix.lower()
    try:
        if ext == ".pdf":
            from pypdf import PdfReader

            return "\n".join((pg.extract_text() or "") for pg in PdfReader(str(path)).pages)
        if ext == ".docx":
            import docx

            return "\n".join(p.text for p in docx.Document(str(path)).paragraphs)
        if ext in (".txt", ".md", ".csv", ".tsv"):
            return path.read_text(encoding="utf-8", errors="ignore")
    except Exception:
        return ""
    return ""


def parse_curation_json(text: str) -> dict | None:
    """Lenient parse of the curation model's JSON (tolerates fences/prose)."""
    if not text:
        return None
    t = text.strip()
    if "```" in t:
        for part in t.split("```"):
            p = part.strip()
            if p.startswith("json"):
                p = p[4:].strip()
            if p.startswith("{"):
                t = p
                break
    start, end = t.find("{"), t.rfind("}")
    if start == -1 or end <= start:
        return None
    try:
        return json.loads(t[start : end + 1])
    except json.JSONDecodeError:
        return None


def build_user_prompt(doc_text: str) -> str:
    clipped = doc_text[:MAX_DOC_CHARS]
    if len(doc_text) > MAX_DOC_CHARS:
        clipped += "\n\n[TRUNCATED]"
    return f"DOCUMENT:\n\n{clipped}"


def to_sft_record(input_text: str, output_text: str) -> dict:
    """Standard chat-format SFT record."""
    return {
        "messages": [
            {"role": "system", "content": advisory_system()},
            {"role": "user", "content": input_text},
            {"role": "assistant", "content": output_text},
        ]
    }


def split_heldout(pairs: list[dict], frac: float, seed: int) -> tuple[list, list]:
    """Stratified by doc_type so the held-out set has both memos + void analyses.
    Held-out NEVER enters training (Gap 3) and seeds new advisory eval cases."""
    rng = random.Random(seed)
    by_type: dict[str, list] = {}
    for p in pairs:
        by_type.setdefault(p.get("doc_type", "unknown"), []).append(p)
    train, heldout = [], []
    for group in by_type.values():
        rng.shuffle(group)
        n = round(len(group) * frac) if len(group) > 1 else 0
        heldout.extend(group[:n])
        train.extend(group[n:])
    return train, heldout


def _write_report(out: Path, kept: list[dict], rejects: list[tuple], stats: dict,
                  n_train: int, n_heldout: int) -> str:
    lines = [
        "# Curation report", "",
        f"- Documents kept: **{len(kept)}**  (train {n_train} / held-out {n_heldout})",
        f"- Documents rejected: **{len(rejects)}**",
        f"- By type: {stats['by_type']}",
        f"- Quality: {stats['quality']}",
        f"- Had source data (Strategy-2 eligible): {stats['had_source_data']}",
        "", "## Top reject reasons", "",
    ]
    for reason, n in sorted(stats["reject_reasons"].items(), key=lambda x: -x[1])[:10]:
        lines.append(f"- {n}× {reason}")
    lines += ["", "## Sample kept pairs (skim these before training)", ""]
    for p in kept[:3]:
        rec = p["record"]["messages"]
        lines.append(f"### {p['source_file']}  ({p['doc_type']}, {p['quality']})")
        lines.append(f"**input:** {rec[1]['content'][:400]}…")
        lines.append(f"**output:** {rec[2]['content'][:400]}…")
        lines.append("")
    text = "\n".join(lines)
    (out / "report.md").write_text(text)
    return text


async def curate_dir(input_dir: str, out_dir: str, model: str, api_key: str,
                     heldout_frac: float, max_docs: int, seed: int) -> None:
    from app.llm.client import get_or_create_client
    from app.llm.types import LLMChatMessage, LLMChatRequest

    client = get_or_create_client("anthropic", api_key, "")
    files = [p for p in sorted(Path(input_dir).rglob("*")) if p.suffix.lower() in SUPPORTED_EXT]
    if max_docs:
        files = files[:max_docs]
    print(f"Found {len(files)} candidate files in {input_dir}\n")

    kept: list[dict] = []
    rejects: list[tuple] = []
    stats = {"by_type": {}, "quality": {}, "had_source_data": {"true": 0, "false": 0}, "reject_reasons": {}}

    for i, path in enumerate(files, 1):
        text = extract_text(path)
        if len(text.strip()) < 200:
            rejects.append((path.name, "empty / too short / unreadable"))
            print(f"  [{i}/{len(files)}] SKIP {path.name} (unreadable)")
            continue
        try:
            resp = await client.chat(LLMChatRequest(
                system=CURATION_SYSTEM,
                messages=[LLMChatMessage(role="user", content=build_user_prompt(text))],
                model=model, max_tokens=4000,
            ))
            data = parse_curation_json(resp.content)
        except Exception as e:
            rejects.append((path.name, f"curation error: {type(e).__name__}"))
            print(f"  [{i}/{len(files)}] ERR  {path.name}: {e}")
            continue
        if not data:
            rejects.append((path.name, "unparseable curation JSON"))
            continue
        if not data.get("usable"):
            reason = (data.get("reject_reason") or "unspecified")[:60]
            rejects.append((path.name, reason))
            stats["reject_reasons"][reason] = stats["reject_reasons"].get(reason, 0) + 1
            print(f"  [{i}/{len(files)}] DROP {path.name} ({reason})")
            continue
        inp, out_text = (data.get("input") or "").strip(), (data.get("output") or "").strip()
        if not inp or not out_text:
            rejects.append((path.name, "missing input/output"))
            continue
        dt = data.get("doc_type", "unknown")
        kept.append({
            "doc_type": dt, "quality": data.get("quality", "?"), "source_file": path.name,
            "had_source_data": bool(data.get("had_source_data")),
            "pii_redacted": data.get("pii_redacted", []),
            "record": to_sft_record(inp, out_text),
        })
        stats["by_type"][dt] = stats["by_type"].get(dt, 0) + 1
        stats["quality"][data.get("quality", "?")] = stats["quality"].get(data.get("quality", "?"), 0) + 1
        stats["had_source_data"]["true" if data.get("had_source_data") else "false"] += 1
        print(f"  [{i}/{len(files)}] KEEP {path.name} ({dt}, {data.get('quality')})")

    await client.aclose()

    train, heldout = split_heldout(kept, heldout_frac, seed)
    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    (out / "train.jsonl").write_text("\n".join(json.dumps(p["record"]) for p in train))
    (out / "heldout.jsonl").write_text("\n".join(json.dumps(p["record"]) for p in heldout))
    report = _write_report(out, kept, rejects, stats, len(train), len(heldout))
    print("\n" + report)
    print(f"\nWrote {out}/train.jsonl, heldout.jsonl, report.md")


def _selftest() -> None:
    import tempfile

    # parse tolerates fences
    p = parse_curation_json('```json\n{"usable": true, "doc_type": "investment_memo"}\n```')
    assert p and p["usable"] is True and p["doc_type"] == "investment_memo"
    assert parse_curation_json("garbage") is None

    # sft record shape
    rec = to_sft_record("the input", "the output")["messages"]
    assert [m["role"] for m in rec] == ["system", "user", "assistant"]
    assert rec[1]["content"] == "the input" and rec[2]["content"] == "the output"

    # stratified split keeps held-out per type, never overlapping train
    pairs = [{"doc_type": "investment_memo", "record": to_sft_record(f"i{i}", f"o{i}")} for i in range(10)]
    pairs += [{"doc_type": "void_analysis", "record": to_sft_record(f"v{i}", f"o{i}")} for i in range(10)]
    tr, ho = split_heldout(pairs, 0.2, seed=1)
    assert len(tr) + len(ho) == 20 and 1 <= len(ho) <= 8
    assert all(h not in tr for h in ho)

    # extraction reads plain text + csv
    with tempfile.TemporaryDirectory() as d:
        f = Path(d) / "memo.txt"
        f.write_text("A void analysis recommending the broker pursue the property.")
        assert "void analysis" in extract_text(f)
        c = Path(d) / "data.csv"
        c.write_text("a,b\n1,2\n")
        assert "a,b" in extract_text(c)

    print("curate.py self-check: OK")


def main() -> int:
    ap = argparse.ArgumentParser(description="Curate raw broker docs into an SFT dataset.")
    ap.add_argument("--selftest", action="store_true", help="run offline structural checks and exit")
    ap.add_argument("--input-dir", help="directory of raw documents")
    ap.add_argument("--out-dir", default="finetune/dataset")
    ap.add_argument("--model", default=os.environ.get("CURATE_MODEL", "claude-sonnet-4-6"))
    ap.add_argument("--heldout-frac", type=float, default=0.15)
    ap.add_argument("--max-docs", type=int, default=0, help="0 = all")
    ap.add_argument("--seed", type=int, default=13)
    args = ap.parse_args()

    if args.selftest:
        _selftest()
        return 0
    if not args.input_dir:
        ap.error("--input-dir is required (or use --selftest)")
    api_key = os.environ.get("CURATE_API_KEY") or os.environ.get("ANTHROPIC_API_KEY", "")
    if not api_key:
        ap.error("set CURATE_API_KEY (or ANTHROPIC_API_KEY)")
    asyncio.run(curate_dir(args.input_dir, args.out_dir, args.model, api_key,
                           args.heldout_frac, args.max_docs, args.seed))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
